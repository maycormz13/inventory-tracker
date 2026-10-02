"""Convert the Excel tracker into data/data.json.
Usage:  pip install openpyxl   then   python convert.py MyTracker.xlsx
Only the INPUT columns are read; formulas (IDs, Days Open, Dashboard) are recalculated by the website."""
import sys, json, datetime as dt
from openpyxl import load_workbook

def clean(v):
    if isinstance(v, (dt.datetime, dt.date)): return v.strftime("%Y-%m-%d")
    if v in ("☑",): return True
    if v in ("☐", None): return ""
    if isinstance(v, str) and v.startswith("="): return ""
    return v

# sheet -> (json key, [(column index starting at 0, field name), ...])
MAP = {
 "INVESTIGATIONS": ("investigations", [(1,"date"),(2,"facility"),(3,"area"),(4,"location"),(5,"sku"),(6,"variance"),(7,"rootCause"),(8,"action"),(9,"owner"),(11,"status"),(12,"due"),(14,"closed"),(15,"evidence")]),
 "ACTIONS": ("actions", [(1,"date"),(2,"facility"),(3,"issue"),(4,"rootCause"),(5,"action"),(6,"owner"),(7,"priority"),(8,"target"),(9,"status"),(10,"closed"),(12,"evidence"),(13,"verified")]),
 "COUNT COMPLETION": ("counts", [(0,"date"),(1,"facility"),(2,"area"),(3,"required"),(4,"completed"),(6,"variances"),(7,"invs"),(8,"owner"),(9,"verified"),(10,"evidence")]),
 "MONTHLY KPI": ("kpis", [(0,"month"),(1,"facility"),(2,"invAcc"),(3,"locAcc"),(4,"countComp"),(5,"fifo"),(6,"invComp"),(7,"verifiedBy"),(8,"evidence")]),
 "ROLLOUT STATUS": ("rollout", [(0,"facility"),(1,"sponsor"),(2,"manager"),(3,"owner"),(4,"training"),(5,"baseline"),(6,"pilotStart"),(7,"pilotEnd"),(8,"approval"),(9,"status"),(10,"notes")]),
 "SETTINGS": ("settings", [(0,"metric"),(1,"target"),(2,"unit"),(3,"owner")]),
}
wb = load_workbook(sys.argv[1], data_only=True)
out = {"updated": dt.date.today().isoformat()}
for sheet, (key, cols) in MAP.items():
    recs = []
    for row in wb[sheet].iter_rows(min_row=5, values_only=True):
        rec = {f: clean(row[i]) if i < len(row) else "" for i, f in cols}
        first = rec[cols[0][1]]
        if first in ("", "Metric", "Facility", "Date", "Month") or first is False: continue
        if key == "kpis" and isinstance(rec["month"], str) and len(rec["month"]) > 7: rec["month"] = rec["month"][:7]
        if any(v not in ("", False) for v in rec.values()): recs.append(rec)
    out[key] = recs
json.dump(out, open("data/data.json", "w"), indent=2, ensure_ascii=False)
print({k: len(v) for k, v in out.items() if isinstance(v, list)})
